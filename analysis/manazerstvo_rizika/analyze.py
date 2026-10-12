"""Read-only extraction and reproducible qualitative risk analysis."""
import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def band(score):
    return 'M' if score <= 2 else 'S' if score <= 4 else 'V'


def analyze(excel, word):
    workbook = load_workbook(excel, data_only=False)
    cached = load_workbook(excel, data_only=True)
    sheet = workbook['MMSP']
    risks, issues = [], []
    for row in range(9, 24):
        values = [sheet.cell(row, col).value for col in range(1, 16)]
        number, classification, title, cause, consequence, p, n, formula, _, acceptance, *details = values
        if type(p) is not int or type(n) is not int or p not in (1, 2, 3) or n not in (1, 2, 3):
            raise ValueError(f'Invalid P/N at row {row}')
        if number != row - 8 or not all(values[c] for c in (1, 2, 3, 4, 9, 10, 11, 12, 13, 14)):
            raise ValueError(f'Incomplete record at row {row}')
        expected = f'F{row}*G{row}'
        if not isinstance(formula, str) or expected not in formula:
            issues.append(f'H{row}: inspect formula')
        score = p * n
        if acceptance not in ('A', 'N'):
            raise ValueError(f'Invalid acceptance at row {row}')
        for col, expected_value in ((8, score), (9, band(score))):
            value = cached['MMSP'].cell(row, col).value
            if value is not None and value != expected_value:
                issues.append(f'Cached value mismatch at row {row}, column {col}')
        risks.append(dict(id=f'R{number:02}', classification=classification, title=title,
                          cause=cause, consequence=consequence, p=p, n=n, score=score,
                          band=band(score), acceptance=acceptance,
                          **dict(zip(('measures','cost_basis','funding','responsibility','monitoring'), details)),
                          source_row=row))
    counts = dict(Counter(r['band'] for r in risks))
    text = '\n'.join(p.text for p in Document(word).paragraphs)
    checks = {'15_risks': len(risks) == 15, 'distribution': counts == {'S':11,'M':3,'V':1},
              'word_distribution': 'Stredných rizík je jedenásť a malých tri' in text,
              'word_top_risk': 'P = 3, N = 2, R = 6' in text,
              'word_scale': '1–2 malé, 3–4 stredné a 6–9 veľké' in text}
    if not all(checks.values()) or issues:
        raise ValueError({'checks': checks, 'issues': issues})
    for r in risks:
        r['rank'] = 1 + sum(other['score'] > r['score'] for other in risks)
        r['sensitivity'] = [dict(p=p,n=n,score=p*n,band=band(p*n),delta=p*n-r['score'])
                            for p,n in ((r['p']-1,r['n']),(r['p']+1,r['n']),
                                        (r['p'],r['n']-1),(r['p'],r['n']+1))
                            if p in (1,2,3) and n in (1,2,3)]
    # Assumptions are illustrative, never inferred from ordinal P or N.
    base = dict(hourly_quote_eur=65, rate_source="Orientačná sadzba uvedená používateľom; nie údaj zo školských dokumentov", other_costs_eur=1500, hours=100,
                nights=10, night_price_eur=60, transport_eur=300)
    base['revenue_eur'] = base['hourly_quote_eur'] * base['hours']
    base_cost = base['other_costs_eur'] + base['nights'] * base['night_price_eur'] + base['transport_eur']
    base_result = base['revenue_eur'] - base_cost
    scenarios = []
    for rid, cases in [('R05', [(5,25),(10,50),(20,100)]),
                       ('R10', [(1,100),(3,120),(6,150)]),
                       ('R11', [(70,0,0),(90,2,100),(120,4,200)])]:
        rows=[]
        for i, case in enumerate(cases, 1):
            if rid == 'R05':
                extra_hours, extra_cost = case
                inputs = dict(repair_hours=extra_hours, consumables_eur=extra_cost)
            elif rid == 'R10':
                extra_hours, rate = case
                extra_cost = extra_hours * rate
                inputs = dict(idle_hours=extra_hours, crane_eur_per_hour=rate)
            else:
                price, nights, transport = case
                extra_hours = 0
                extra_cost = (base['nights']+nights)*price-base['nights']*base['night_price_eur']+transport
                inputs = dict(night_price_eur=price, extra_nights=nights, extra_transport_eur=transport)
            result = base_result-extra_cost
            rows.append(dict(case=i,inputs=inputs,extra_cost_eur=extra_cost,result_eur=result,loss_eur=base_result-result,
                             hourly_drop_eur=round(base_result/base["hours"]-result/(base["hours"]+extra_hours),2),
                             work_hours=base['hours'] + (extra_hours if rid == 'R05' else 0),
                             waiting_hours=extra_hours if rid == 'R10' else 0,
                             hours=base['hours']+extra_hours, result_per_hour_eur=round(result/(base['hours']+extra_hours),2)))
        scenarios.append(dict(id=rid,cases=rows))
    return dict(schema_version=2, horizon_months=12, sources=[dict(name=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in (excel,word)],
                audit=checks, counts=counts, risks=sorted(risks,key=lambda r:(-r['score'],-r['n'],r['id'])),
                matrix=[dict(p=p,n=n,score=p*n,band=band(p*n),ids=[r['id'] for r in risks if r['p']==p and r['n']==n]) for p in (3,2,1) for n in (1,2,3)],
                economics=dict(model_only=True,base=base,base_cost_eur=base_cost,base_result_eur=base_result,base_result_per_hour_eur=round(base_result/base['hours'],2),scenarios=scenarios))


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('excel',type=Path)
    parser.add_argument('word',type=Path)
    parser.add_argument('--output',type=Path,default=Path(__file__).with_name('risk_analysis.json'))
    args=parser.parse_args()
    data=analyze(args.excel,args.word)
    args.output.write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    print(json.dumps({'audit':data['audit'],'counts':data['counts']},ensure_ascii=False))
